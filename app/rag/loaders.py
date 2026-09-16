import json
from pathlib import Path
from typing import Callable

from langchain_core.documents import Document

from app.logging_config import get_logger

logger = get_logger(__name__)

# MIME_TO_EXT/LOADER_MAP 为模块级常量，模块加载时构建一次后复用

MIME_TO_EXT: dict[str, str] = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/csv": ".csv",
    "text/html": ".html",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/toml": ".toml",
    "application/json": ".json",
}


def _load_pdf(path: str) -> list[Document]:
    """加载 PDF 文件"""
    from langchain_community.document_loaders import PyPDFLoader

    loader = PyPDFLoader(path)
    return loader.load()


def _load_text(path: str) -> list[Document]:
    """加载文本、Markdown文件"""
    from langchain_community.document_loaders import TextLoader

    loader = TextLoader(path, encoding="utf-8")
    return loader.load()


def _load_docx(path: str) -> list[Document]:
    """加载 Word 文档"""
    from langchain_community.document_loaders import Docx2txtLoader

    loader = Docx2txtLoader(path)
    return loader.load()


def _load_csv(path: str) -> list[Document]:
    """加载 CSV 文件"""
    from langchain_community.document_loaders import CSVLoader

    loader = CSVLoader(file_path=path)
    return loader.load()


def _load_html(path: str) -> list[Document]:
    """加载 HTML 文件"""
    from langchain_community.document_loaders import BSHTMLLoader

    loader = BSHTMLLoader(path)
    return loader.load()


def _load_xlsx(path: str) -> list[Document]:
    """加载 Excel 文件"""
    # from langchain_community.document_loaders import XLSXLoader

    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None)  # dict[sheet_name, DataFrame]
    return [
        Document(
            page_content=f"### Sheet: {name}\n\n{df.astype(str).to_csv(index=False)}",
            metadata={"source": path},
        )
        for name, df in sheets.items()
    ]


def _load_pptx(path: str) -> list[Document]:
    """加载 PowerPoint 文件"""
    # from langchain_community.document_loaders import PPTXLoader

    from pptx import Presentation

    prs = Presentation(path)
    docs = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        docs.append(
            Document(
                page_content=f"### Slide {i}\n\n" + "\n\n".join(parts),
                metadata={"source": path},
            )
        )
    return docs


def _load_toml(path: str) -> list[Document]:
    """加载 TOML 文件"""
    from langchain_community.document_loaders import TOMLLoader

    loader = TOMLLoader(path)
    return loader.load()


def _load_json(path: str) -> list[Document]:
    """自定义 JSON 加载器，兼容多种格式：

    1. 法规 seed_data 格式: {title, articles: [{article_number, chapter?, content}]}  → 每条条文一个 Document

    2. 对象数组: [{key: value, ...}, ...] → 每个对象一个 Document

    3. 字符串数组: ["text", ...] → 每个字符串一个 Document

    4. 普通对象: {key: value, ...} → 整体一个 Document
    """
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # 格式 1：法规 seed_data — 顶层有 articles 数组
    if isinstance(data, dict) and isinstance(data.get("articles"), list):
        # Path(path).stem 会自动去掉文件扩展名，例如 "example.json" 会变成 "example"。
        regulation_title = data.get("title") or data.get("name") or Path(path).stem
        regulation_type = data.get("regulation_type", "")

        docs: list[Document] = []
        for art in data["articles"]:
            content = art.get("content", "").strip()
            if not content:
                continue

            parts = [f"# {regulation_title}"]
            if regulation_type:
                parts.append(f"类型：{regulation_type}")

            if art.get("chapter"):
                parts.append(f"章节：{art['chapter']}")

            if art.get("article_number"):
                parts.append(f"条文编号：{art['article_number']}")

            parts.append("")
            parts.append(content)
            docs.append(
                Document(
                    page_content="\n".join(parts),
                    metadata={
                        "source": path,
                        "regulation_title": regulation_title,
                        "article_number": art.get("article_number", ""),
                    },
                )
            )
        if docs:
            logger.info(
                f"json (regulation format): extracted {len(docs)} articles from {regulation_title}"
            )
            return docs

    # 格式 2：对象数组
    if isinstance(data, list) and data and isinstance(data[0], dict):
        docs = []
        for item in data:
            lines = []
            for k, v in item.items():
                if isinstance(v, (dict, list)):
                    lines.append(f"{k}: {json.dumps(v, ensure_ascii=False)}")
                else:
                    lines.append(f"{k}: {v}")
            docs.append(
                Document(
                    page_content="\n".join(lines),
                    metadata={"source": path},
                )
            )
        logger.info(f"json (list of objects): {len(docs)} items")
        return docs

    # 格式 3：字符串数组
    if isinstance(data, list) and data and isinstance(data[0], str):
        docs = [
            Document(page_content=item, metadata={"source": path})
            for item in data
            if isinstance(item, str) and item.strip()
        ]
        logger.info(f"json (list of strings): {len(docs)} items")
        return docs

    # 格式 4：普通对象或其他 — 整体序列化为文本
    docs = [
        Document(
            page_content=json.dumps(data, ensure_ascii=False, indent=2),
            metadata={"source": path},
        )
    ]
    logger.info("json (generic): single document")
    return docs


# 函数处理映射
LOADER_MAP: dict[str, Callable[[str], list[Document]]] = {
    "application/pdf": _load_pdf,
    "text/markdown": _load_text,
    "text/plain": _load_text,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _load_docx,
    "text/csv": _load_csv,
    "text/html": _load_html,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _load_xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _load_pptx,
    "application/toml": _load_toml,
    "application/json": _load_json,
}

# 两个变量的key值要相同
assert LOADER_MAP.keys() == MIME_TO_EXT.keys(), "LOADER_MAP 和 MIME_TO_EXT 的 key 不一致"


def load_multi_documents(file_path: str, mime_type: str) -> list[Document]:
    """从磁盘加载文件并返回LangChain的Document对象列表。"""
    loader = LOADER_MAP.get(mime_type)
    if loader is None:
        logger.error(f"unsupported MIME type for loading: {mime_type}")
        raise ValueError(f"Unsupported MIME type for loading: {mime_type}")
    return loader(file_path)
